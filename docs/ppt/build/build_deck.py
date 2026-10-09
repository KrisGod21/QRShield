"""Build the 6-slide CodeAstra 2.0 idea deck for QR Shield on the official template art.

Template geometry comes from the PDF (1440 x 810 pt); slide is 13.333 x 7.5 in,
so 1 template pt = 1/108 in.
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
A = os.path.join(HERE, "assets")
OUT = os.path.join(HERE, "..", "QR_Shield_CodeAstra.pptx")

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xB8, 0xC4, 0xDE)
TEAL = RGBColor(0x2D, 0xD4, 0xBF)
RED = RGBColor(0xEF, 0x44, 0x44)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
PANEL = RGBColor(0x14, 0x1E, 0x35)
PANEL2 = RGBColor(0x1C, 0x29, 0x47)
LINE = RGBColor(0x33, 0x44, 0x6B)
BODY = "Segoe UI"
SERIF = "Times New Roman"


def pt(v):  # template points -> EMU
    return Inches(v / 108)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def base(bg, title=None):
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(os.path.join(A, bg), 0, 0, prs.slide_width, prs.slide_height)
    s.shapes.add_picture(os.path.join(A, "header_mct.png"), pt(331.5), 0, pt(777), pt(114))
    s.shapes.add_picture(os.path.join(A, "logo_abit.png"), pt(81), pt(12), pt(102), pt(102))
    s.shapes.add_picture(os.path.join(A, "logo_synergy.png"), pt(1266.5), pt(12), pt(92.2), pt(102))
    if title:
        text(s, title, pt(81), pt(132), Inches(11), Inches(0.6), 26, font=SERIF)
    return s


def text(s, t, x, y, w, h, size=14, color=WHITE, bold=False, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = ln if isinstance(ln, list) else [(ln, {})]
        for seg, fmt in runs:
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = fmt.get("font", font)
            f.size = Pt(fmt.get("size", size))
            f.bold = fmt.get("bold", bold)
            f.italic = fmt.get("italic", italic)
            f.color.rgb = fmt.get("color", color)
    return tb


def box(s, x, y, w, h, fill=PANEL, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, alpha=None, radius=0.08, lw=1):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if alpha is not None:
            from pptx.oxml.ns import qn
            sf = sh.fill._xPr.find(qn("a:solidFill"))
            clr = sf[0]
            a = clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
            clr.append(a)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def arrow(s, x1, y1, x2, y2, color=TEAL, w=2):
    c = s.shapes.add_connector(1, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(w)
    from pptx.oxml.ns import qn
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


# =========================================================== 1. Team details
s = prs.slides.add_slide(BLANK)
s.shapes.add_picture(os.path.join(A, "bg_title.jpg"), 0, 0, prs.slide_width, prs.slide_height)
s.shapes.add_picture(os.path.join(A, "header_mct.png"), pt(331.5), 0, pt(777), pt(114))
s.shapes.add_picture(os.path.join(A, "logo_abit.png"), pt(81), pt(12), pt(102), pt(102))
s.shapes.add_picture(os.path.join(A, "logo_synergy.png"), pt(1266.5), pt(12), pt(92.2), pt(102))
s.shapes.add_picture(os.path.join(A, "codeastra.png"), pt(352.3), pt(-53.6), pt(670.5), pt(447))
text(s, "Team Details", pt(420), pt(236), pt(530), pt(52), 28, font=SERIF, align=PP_ALIGN.CENTER)
fields = [("Team ID : ", "<TEAM ID>"), ("Team Name : ", "<TEAM NAME>"), ("Leader Name : ", "<LEADER NAME>"),
          ("Domain Chosen : ", "Cybersecurity"),
          ("Problem Statement : ", "QR Code Safety Scanner – scans QR codes before opening them and warns users if the destination looks malicious")]
y = 352
for label, val in fields:
    h = 110 if label.startswith("Problem") else 46
    text(s, [[(label, {"font": SERIF, "size": 22}), (val, {"font": SERIF, "size": 20, "color": RGBColor(0xDD, 0xF7, 0xF3)})]],
         pt(132), pt(y), pt(820), pt(h), 22)
    y += 47
# product lockup
box(s, pt(1000), pt(600), pt(380), pt(150), fill=RGBColor(0x0B, 0x12, 0x20), line=TEAL, alpha=0.72, radius=0.12, lw=1.5)
text(s, [[("🛡 ", {"size": 30}), ("QR Shield", {"size": 30, "bold": True, "color": TEAL})]],
     pt(1015), pt(610), pt(360), pt(60), align=PP_ALIGN.CENTER)
text(s, "Know what a QR will make you do — before you do it.", pt(1015), pt(672), pt(350), pt(70), 15,
     color=WHITE, align=PP_ALIGN.CENTER, italic=True)

# =========================================================== 2. Problem analysis
s = base("bg_content.jpg", "Problem Analysis")
stats = [("24 bn", "UPI transactions every month", "NPCI, Sep 2026", TEAL),
         ("73 crore", "UPI QR codes in circulation", "Worldline, 2025", TEAL),
         ("2.7×", "QR-code fraud cases in 2 years\n14,625 → 39,638", "Lok Sabha, FY22→FY24", RED),
         ("₹805 cr", "lost to 10.6 lakh UPI frauds\nin just 8 months", "MoF, Lok Sabha, Apr–Nov 2025", RED)]
x0, wtile, gap = 0.72, 2.86, 0.2
for i, (num, cap, src, col) in enumerate(stats):
    x = Inches(x0 + i * (wtile + gap))
    box(s, x, Inches(1.95), Inches(wtile), Inches(1.55), fill=PANEL, line=col, alpha=0.85, lw=1.5)
    text(s, num, x, Inches(2.0), Inches(wtile), Inches(0.62), 32, color=col, bold=True, align=PP_ALIGN.CENTER)
    text(s, cap, x + Inches(0.1), Inches(2.6), Inches(wtile - 0.2), Inches(0.62), 12.5, align=PP_ALIGN.CENTER)
    text(s, src, x, Inches(3.18), Inches(wtile), Inches(0.28), 9.5, color=MUTED, align=PP_ALIGN.CENTER, italic=True)

text(s, "HOW PEOPLE GET CHEATED", Inches(0.75), Inches(3.68), Inches(6), Inches(0.3), 11, color=MUTED, bold=True)
pats = [("💸", "“Scan to receive money”", "Fake buyers & ‘customer care’ send a QR. Scanning a UPI QR always PAYS — victims enter their PIN."),
        ("🏪", "Swapped shop stickers", "Fraudster’s QR pasted over the shop’s: Khajuraho 2025, Aizawl 2024, Goa ₹49.68 lakh (2026)."),
        ("🎣", "Fake bank / KYC / e-challan", "Lookalike sites, AutoPay mandates, APK downloads and *401# call-forward codes.")]
for i, (ic, ttl, body) in enumerate(pats):
    x = Inches(0.72 + i * 4.0)
    box(s, x, Inches(4.0), Inches(3.82), Inches(1.45), fill=PANEL2, line=None, alpha=0.85)
    text(s, ic, x + Inches(0.12), Inches(4.1), Inches(0.6), Inches(0.6), 24)
    text(s, ttl, x + Inches(0.72), Inches(4.1), Inches(3.0), Inches(0.4), 14, bold=True)
    text(s, body, x + Inches(0.72), Inches(4.5), Inches(3.0), Inches(0.95), 11, color=MUTED)

box(s, Inches(0.72), Inches(5.65), Inches(11.88), Inches(1.45), fill=RGBColor(0x2A, 0x10, 0x16), line=RED, alpha=0.8, lw=1.25)
text(s, [[("Root cause  ", {"bold": True, "color": RED, "size": 13}),
          ("NPCI ended P2P ‘collect’ requests on 1 Oct 2025 — so the QR is now the scammer’s main lever. "
           "RBI, NPCI and police all say “don’t scan unknown QRs”, but nobody gives people a tool to check one.", {"size": 13})],
         [("The gap  ", {"bold": True, "color": AMBER, "size": 13}),
          ("Existing scanners (Kaspersky, Trend Micro, phone cameras) only match links against blacklists. "
           "None understand UPI: who you pay, how much, AutoPay — or explain why in Hindi or Marathi.", {"size": 13})]],
     Inches(0.95), Inches(5.72), Inches(11.5), Inches(1.35), 13, anchor=MSO_ANCHOR.MIDDLE)

# =========================================================== 3. Solution & features
s = base("bg_content.jpg", "Proposed Solution & Key Features")
# phone mock
px, py, pw, ph = Inches(0.8), Inches(1.9), Inches(3.55), Inches(5.35)
box(s, px, py, pw, ph, fill=RGBColor(0x0B, 0x12, 0x20), line=RGBColor(0x55, 0x66, 0x88), radius=0.09, lw=2)
ix, iw = px + Inches(0.18), pw - Inches(0.36)
box(s, ix, py + Inches(0.22), iw, Inches(0.62), fill=RED, line=None, radius=0.2)
text(s, [[("✕  DANGER — do not open", {"bold": True, "size": 14})], [("Risk 100/100 · 0.4 s", {"size": 9, "color": RGBColor(0xFF, 0xE0, 0xE0)})]],
     ix + Inches(0.12), py + Inches(0.24), iw, Inches(0.6), 14)
text(s, "You are about to PAY ₹4,999 to “Amazon Refund Dept”", ix, py + Inches(0.92), iw, Inches(0.55), 12.5, bold=True)
box(s, ix, py + Inches(1.5), iw, Inches(1.02), fill=PANEL, line=LINE, radius=0.12)
text(s, [[("YOU PAY  ", {"bold": True, "color": RED, "size": 10}), ("₹4,999 🔒", {"bold": True, "size": 20})],
         [("refund.desk2291@ybl", {"size": 10, "color": TEAL, "font": "Consolas"})]],
     ix + Inches(0.1), py + Inches(1.52), iw, Inches(0.62), 12)
box(s, ix, py + Inches(2.17), iw, Inches(0.35), fill=RGBColor(0x45, 0x16, 0x1C), line=None, shape=MSO_SHAPE.RECTANGLE)
text(s, "💸 Scanning a QR never gives you money", ix + Inches(0.08), py + Inches(2.18), iw, Inches(0.33), 10, bold=True,
     color=RGBColor(0xFE, 0xCA, 0xCA))
reasons = ["⛔ Says ‘refund’ — you never pay to get a refund", "⚠️ Official name on an ordinary app UPI ID",
           "⚠️ Amount locked; receiver looks like a person", "👥 Reported 6 times by other users"]
for i, rsn in enumerate(reasons):
    yy = py + Inches(2.62 + i * 0.47)
    box(s, ix, yy, iw, Inches(0.41), fill=PANEL, line=RGBColor(0x5A, 0x24, 0x2A), radius=0.18)
    text(s, rsn, ix + Inches(0.08), yy + Inches(0.03), iw - Inches(0.1), Inches(0.38), 9.5, anchor=MSO_ANCHOR.MIDDLE)
by = py + ph - Inches(0.62)
box(s, ix, by, iw / 2 - Inches(0.05), Inches(0.42), fill=TEAL, line=None, radius=0.25)
text(s, "🛡 Don't pay", ix, by + Inches(0.04), iw / 2, Inches(0.36), 11, bold=True, color=RGBColor(4, 18, 26), align=PP_ALIGN.CENTER)
box(s, ix + iw / 2 + Inches(0.05), by, iw / 2 - Inches(0.05), Inches(0.42), fill=None, line=LINE, radius=0.25)
text(s, "Open anyway", ix + iw / 2, by + Inches(0.04), iw / 2, Inches(0.36), 11, color=MUTED, align=PP_ALIGN.CENTER)
text(s, "EN | हिं | मरा", px, py + ph + Inches(0.0), pw, Inches(0.25), 9, color=MUTED, align=PP_ALIGN.CENTER)

feats = [("🔍", "UPI-aware decoding", "Shows who you PAY, how much, AutoPay mandates; flags lookalike UPI IDs (@okaxisbank, @ybI) & fake ‘refund’ / ‘KYC’ notes"),
         ("🙋", "“Receiving money?” mode", "One tap tells us your intent. Any UPI QR then becomes an instant red stop — kills the #1 scam"),
         ("🗣", "Reasons in your language", "Plain explanations in English, हिंदी & मराठी — written for shopkeepers & parents, not engineers"),
         ("🏪", "Shopkeeper Mode", "Register your QR once; a 2-second scan of the counter sticker catches a swap before customers pay"),
         ("🧩", "Every QR type", "Links, Wi-Fi (fake ‘RailWire’), *401# call-forward codes, APKs, AutoPay, Bharat QR checksum"),
         ("👥", "Community + offline", "Hashed ‘Report this QR’ counts; built-in rules still give a verdict with no internet")]
fx0, fy0, fw, fh = 4.65, 1.9, 4.05, 1.68
for i, (ic, ttl, body) in enumerate(feats):
    x = Inches(fx0 + (i % 2) * (fw + 0.12))
    y = Inches(fy0 + (i // 2) * (fh + 0.12))
    box(s, x, y, Inches(fw), Inches(fh), fill=PANEL, line=LINE, alpha=0.88)
    text(s, ic, x + Inches(0.12), y + Inches(0.12), Inches(0.6), Inches(0.6), 24)
    text(s, ttl, x + Inches(0.75), y + Inches(0.14), Inches(fw - 0.85), Inches(0.4), 14.5, bold=True, color=TEAL)
    text(s, body, x + Inches(0.75), y + Inches(0.55), Inches(fw - 0.85), Inches(1.1), 11)

# =========================================================== 4. Technical approach
s = base("bg_content.jpg", "Technical approach & Innovation")
flow = [("📱 Phone PWA", "Camera / gallery decode\nZXing-C++ WASM\nsends only decoded text"),
        ("🧭 Classifier", "UPI · Bharat QR · URL\nWi-Fi · tel/SMS · text\nAutoPay · app links"),
        ("🔬 Analyzers", "UPI semantics & handles\nURL lookalikes & intel\nWi-Fi / USSD rules"),
        ("⚖️ Risk engine", "Weighted signals → 0–100\nSAFE / CAUTION / DANGER\n+ reasons in 3 languages"),
        ("✅ Verdict < 3 s", "Don’t open / Open anyway\nReport · History\nShopkeeper check")]
bw, bh, by0 = 2.2, 1.45, 1.95
for i, (ttl, body) in enumerate(flow):
    x = Inches(0.72 + i * (bw + 0.22))
    col = TEAL if i in (0, 4) else LINE
    box(s, x, Inches(by0), Inches(bw), Inches(bh), fill=PANEL, line=col, alpha=0.9, lw=1.5)
    text(s, ttl, x, Inches(by0 + 0.08), Inches(bw), Inches(0.4), 13.5, bold=True, align=PP_ALIGN.CENTER)
    text(s, body, x + Inches(0.06), Inches(by0 + 0.5), Inches(bw - 0.12), Inches(0.95), 10, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow(s, x + Inches(bw), Inches(by0 + bh / 2), x + Inches(bw + 0.22), Inches(by0 + bh / 2))
text(s, "FastAPI server (Python) — the link is never opened on the phone", Inches(3.1), Inches(1.62), Inches(7.3), Inches(0.3),
     10.5, color=TEAL, italic=True, align=PP_ALIGN.CENTER)

side = [("🧠 Offline rules (always on)", ["Unicode TR39 ‘skeleton’ → lookalike letters (onlinesbі = onlinesbi)",
                                          "Typo-squats of 50+ Indian banks, govt & brands",
                                          "181 real UPI handles · scam words EN/HI/MR",
                                          ".bank.in / .gov.in trust · risky TLDs · APK/USSD"]),
        ("🌐 Online checks (time-boxed ≤ 2.2 s)", ["SSRF-safe redirect expansion (IP-pinned, 5 hops)",
                                                 "RDAP domain age (new = risky)",
                                                 "OpenPhish · PhishTank · URLhaus feeds",
                                                 "Google Safe Browsing (optional key)"])]
for i, (ttl, items) in enumerate(side):
    x = Inches(0.72 + i * 4.1)
    box(s, x, Inches(3.65), Inches(3.95), Inches(2.0), fill=PANEL2, line=None, alpha=0.88)
    text(s, ttl, x + Inches(0.15), Inches(3.72), Inches(3.7), Inches(0.38), 13, bold=True, color=TEAL if i == 0 else AMBER)
    text(s, ["•  " + it for it in items], x + Inches(0.15), Inches(4.12), Inches(3.7), Inches(1.5), 10.5)

box(s, Inches(8.95), Inches(3.65), Inches(3.65), Inches(2.0), fill=RGBColor(0x0E, 0x2B, 0x2A), line=TEAL, alpha=0.9, lw=1.25)
text(s, "💡 Innovation / USP", Inches(9.1), Inches(3.72), Inches(3.4), Inches(0.38), 13, bold=True, color=TEAL)
text(s, ["•  First scanner that explains UPI before the payment app opens",
         "•  Explainable & deterministic — no black box; works offline",
         "•  Catches brand-new scams blacklists haven’t seen",
         "•  Merchant-side sticker verification"], Inches(9.1), Inches(4.12), Inches(3.4), Inches(1.5), 10.5)

text(s, "TECH STACK", Inches(0.75), Inches(5.85), Inches(3), Inches(0.3), 11, color=MUTED, bold=True)
stack = ["React + TypeScript", "Vite PWA", "barcode-detector (ZXing WASM)", "Python FastAPI", "httpx", "tldextract + PSL",
         "RapidFuzz", "Unicode TR39", "SQLite", "pytest (140+ tests)"]
x = 0.75
yrow = 6.2
for item in stack:
    w = 0.22 + 0.085 * len(item)
    if x + w > 12.6:
        x, yrow = 0.75, yrow + 0.48
    box(s, Inches(x), Inches(yrow), Inches(w), Inches(0.38), fill=PANEL, line=TEAL, radius=0.5)
    text(s, item, Inches(x), Inches(yrow + 0.04), Inches(w), Inches(0.32), 10.5, align=PP_ALIGN.CENTER)
    x += w + 0.12

# =========================================================== 5. Impact & future scope
s = base("bg_content.jpg", "Impact & Future Scope")
text(s, "TARGET USERS", Inches(0.75), Inches(1.9), Inches(4), Inches(0.3), 11, color=MUTED, bold=True)
users = [("👵", "UPI users & senior citizens", "Scan before you pay"), ("🏪", "Kirana & small shopkeepers", "Daily sticker check"),
         ("🛍", "OLX/Quikr sellers & students", "‘Scan to receive’ targets"), ("👮", "Banks, police & colleges", "Awareness drives")]
for i, (ic, ttl, sub) in enumerate(users):
    x = Inches(0.72 + i * 3.0)
    box(s, x, Inches(2.22), Inches(2.85), Inches(1.15), fill=PANEL, line=LINE, alpha=0.88)
    text(s, ic, x + Inches(0.1), Inches(2.33), Inches(0.6), Inches(0.6), 26)
    text(s, ttl, x + Inches(0.72), Inches(2.3), Inches(2.05), Inches(0.6), 12.5, bold=True)
    text(s, sub, x + Inches(0.72), Inches(2.88), Inches(2.05), Inches(0.4), 10.5, color=MUTED)

text(s, "IMPACT", Inches(0.75), Inches(3.6), Inches(4), Inches(0.3), 11, color=MUTED, bold=True)
impact = [("🛑", "Stops the #1 UPI scam at the moment of scanning — before the PIN"),
          ("🗣", "Non-technical users understand WHY, in their own language"),
          ("🏪", "Shopkeepers catch a swapped sticker before losing a day’s sales"),
          ("🔒", "Privacy-first: reports stored as hashes; links never open on the phone")]
for i, (ic, t_) in enumerate(impact):
    x = Inches(0.72 + (i % 2) * 6.0)
    y = Inches(3.92 + (i // 2) * 0.62)
    box(s, x, y, Inches(5.85), Inches(0.54), fill=PANEL2, line=None, alpha=0.85, radius=0.2)
    text(s, [[(ic + "  ", {"size": 14}), (t_, {"size": 12})]], x + Inches(0.12), y + Inches(0.05), Inches(5.65), Inches(0.46),
         12, anchor=MSO_ANCHOR.MIDDLE)

text(s, "FUTURE SCOPE — ROADMAP", Inches(0.75), Inches(5.3), Inches(5), Inches(0.3), 11, color=MUTED, bold=True)
road = [("SDK / API", "for UPI apps & banks"), ("Label OCR", "printed name vs QR payload"),
        ("1930 & Chakshu", "one-tap reporting; DoT FRI"), ("More languages", "+ voice explanations"),
        ("On-device mode", "fully offline rules")]
ly = Inches(6.05)
arrow(s, Inches(0.8), ly, Inches(12.55), ly, color=TEAL, w=2.5)
for i, (ttl, sub) in enumerate(road):
    cx = 1.45 + i * 2.55
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.11), ly - Inches(0.11), Inches(0.22), Inches(0.22))
    dot.fill.solid(); dot.fill.fore_color.rgb = TEAL; dot.line.fill.background()
    text(s, ttl, Inches(cx - 1.2), ly + Inches(0.18), Inches(2.4), Inches(0.35), 12.5, bold=True, align=PP_ALIGN.CENTER)
    text(s, sub, Inches(cx - 1.2), ly + Inches(0.5), Inches(2.4), Inches(0.35), 10, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, f"Phase {i + 1}", Inches(cx - 1.0), ly - Inches(0.45), Inches(2.0), Inches(0.3), 9.5, color=TEAL, align=PP_ALIGN.CENTER)

# =========================================================== 6. Supporting information
s = base("bg_content.jpg", "Supporting Information")
text(s, "USER WORKFLOW", Inches(0.75), Inches(1.9), Inches(4), Inches(0.3), 11, color=MUTED, bold=True)
steps = ["1  Scan / upload", "2  Decode on phone", "3  Analyze on server", "4  Verdict + reasons", "5  Don’t open · Report"]
for i, st in enumerate(steps):
    x = Inches(0.72 + i * 2.4)
    sh = s.shapes.add_shape(MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, x, Inches(2.25), Inches(2.45), Inches(0.62))
    sh.fill.solid(); sh.fill.fore_color.rgb = TEAL if i == 4 else PANEL2; sh.line.fill.background()
    text(s, st, x + Inches(0.3), Inches(2.33), Inches(1.95), Inches(0.46), 11.5, bold=True,
         color=RGBColor(4, 18, 26) if i == 4 else WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

box(s, Inches(0.72), Inches(3.15), Inches(5.0), Inches(2.75), fill=PANEL, line=GREEN, alpha=0.88, lw=1.25)
text(s, "✅ Built & validated (working prototype)", Inches(0.9), Inches(3.25), Inches(4.7), Inches(0.38), 13.5, bold=True, color=GREEN)
text(s, ["•  140+ automated tests: UPI parser, homoglyphs, scoring, SSRF guard",
         "•  14 printable demo QRs — every one gives the expected verdict",
         "•  SSRF guard tested vs private IPs, cloud metadata, DNS rebinding, redirect loops",
         "•  Verdict in ~0.1 s offline, < 3 s with live checks",
         "•  Demo mode: reserved .test domains, works with no internet",
         "•  Installable PWA: Android & iPhone browsers"],
     Inches(0.9), Inches(3.68), Inches(4.75), Inches(2.2), 11.5)

box(s, Inches(5.95), Inches(3.15), Inches(6.65), Inches(2.75), fill=PANEL, line=LINE, alpha=0.88)
text(s, "📚 References", Inches(6.12), Inches(3.25), Inches(6.3), Inches(0.38), 13.5, bold=True, color=TEAL)
refs = ["NPCI — UPI Linking Specification v1.6 (upi://pay parameters, signed QR)",
        "RBI consumer advisory, 28 Jan 2022 — “receipt of money does not require scanning QR”",
        "Lok Sabha USQ 2245, 9 Dec 2024 — QR-code fraud FY22–FY25",
        "MoF Lok Sabha reply, Dec 2025 — UPI fraud ₹805 cr (Apr–Nov 2025)",
        "NPCI circular, 29 Jul 2025 — P2P collect discontinued from 1 Oct 2025",
        "RBI circular, Apr 2025 — exclusive .bank.in domain for banks",
        "DoT advisory, Jan 2024 — *401# call-forwarding scam",
        "Unicode TR39 confusables · OWASP SSRF Prevention Cheat Sheet",
        "Interisle Phishing Landscape 2025 · Unit 42 newly-registered-domain study"]
text(s, ["•  " + r for r in refs], Inches(6.12), Inches(3.66), Inches(6.35), Inches(2.2), 10.5)

text(s, "LIVE DEMO SCENARIOS — every QR on our printed sheet gives the expected verdict", Inches(0.75), Inches(6.05),
     Inches(11), Inches(0.3), 11, color=MUTED, bold=True)
demo = [("SBI .bank.in", "SAFE"), ("Kirana UPI", "SAFE"), ("Home Wi-Fi", "SAFE"), ("Café open Wi-Fi", "CAUTION"),
        ("‘Refund’ UPI", "DANGER"), ("@okaxisbank KYC", "DANGER"), ("AutoPay trap", "DANGER"), ("onlinesbl.test", "DANGER"),
        ("onlinesbі (Cyrillic)", "DANGER"), ("Short link → fake HDFC", "DANGER"), ("e-challan APK", "DANGER"),
        ("Fake RailWire", "DANGER"), ("*401# code", "DANGER"), ("Swapped sticker", "SHOP ✕")]
cols = {"SAFE": GREEN, "CAUTION": AMBER, "DANGER": RED, "SHOP ✕": RED}
x, yrow = 0.75, 6.4
for name, v in demo:
    w = 0.3 + 0.075 * len(name)
    if x + w > 12.65:
        x, yrow = 0.75, yrow + 0.42
    box(s, Inches(x), Inches(yrow), Inches(w), Inches(0.34), fill=PANEL2, line=cols[v], radius=0.5)
    text(s, name, Inches(x), Inches(yrow + 0.03), Inches(w), Inches(0.3), 9.5, align=PP_ALIGN.CENTER, color=WHITE)
    x += w + 0.08

prs.save(OUT)
print("saved", os.path.abspath(OUT))
